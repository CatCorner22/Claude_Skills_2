#!/usr/bin/env python3
"""Audit a .pptx against the assertion-evidence checklist.

Usage:
  python ae_lint.py deck.pptx [--json] [--max-body-words N]
  python ae_lint.py --self-test    # lint a deck built here that exercises each check

Reports headline violations, bullet characters, body word-count overruns, banned
formatting (italics, underline, shouting caps), missing source tags on slides
that carry figures, and unexpected fonts. The formatting and font checks cover
the headline and the body; the bottom zone (source tags, slide numbers) is
exempt, since reference listings are conventionally set apart. Exits non-zero
when it finds an ERROR (warnings do not fail the run). Fix the generator and
rebuild rather than hand-patching the packed XML.

Requires: python-pptx  (pip install python-pptx)
"""
from __future__ import annotations

import argparse
import json
import re
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    sys.exit("python-pptx is required: pip install python-pptx")

HEAD_LINE_CHARS = 52          # ~105 chars over two lines at 28 pt / 11.52 in
MAX_BODY_WORDS = 40
BOTTOM_ZONE_IN = 6.5          # shapes below this top are source tags / slide numbers
ALLOWED_FONTS = {"calibri", "arial", "gotham", "montserrat", "goudy old style"}
BULLET_CHARS = "•‣▪◦·●■–—"
BULLET_RE = re.compile(r"^\s*[" + re.escape(BULLET_CHARS) + r"]\s+")
DASH_BULLET_RE = re.compile(r"^\s*[-*]\s+")
DIGIT_RE = re.compile(r"\d")
# All-caps words worth flagging (real words shouted), sparing short acronyms.
CAPS_RE = re.compile(r"\b[A-Z]{6,}\b")
CAPS_ALLOW = {"AUTORECON"}    # domain terms that are legitimately upper


def estimate_lines(text, max_chars=HEAD_LINE_CHARS):
    lines, cur = 0, ""
    for word in str(text).split():
        cand = (cur + " " + word).strip()
        if len(cand) > max_chars and cur:
            lines += 1
            cur = word
        else:
            cur = cand
    if cur:
        lines += 1
    return max(lines, 1)


def _emu_to_in(v):
    return (v or 0) / 914400.0


class Findings:
    def __init__(self):
        self.items = []

    def add(self, slide, level, check, msg):
        self.items.append({"slide": slide, "level": level, "check": check, "message": msg})

    @property
    def errors(self):
        return [f for f in self.items if f["level"] == "error"]


def _text_shapes(slide):
    """Yield (shape, top_in) for shapes with a text frame, non-empty."""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            yield sh, _emu_to_in(sh.top)


def lint_slide(idx, slide, fnd, max_words):
    texts = list(_text_shapes(slide))
    # Headline = highest (smallest top) text shape not in the bottom zone.
    candidates = [(sh, t) for sh, t in texts if t < BOTTOM_ZONE_IN]
    headline_shape = min(candidates, key=lambda p: p[1])[0] if candidates else None
    headline = headline_shape.text_frame.text.strip() if headline_shape else ""

    if not headline:
        fnd.add(idx, "error", "headline", "no headline text found on the slide")
    else:
        lines = estimate_lines(headline)
        if lines > 2 or len(headline) > 110:
            fnd.add(idx, "error", "headline",
                    f"headline exceeds two lines (~{lines}): {headline!r}")
        if len(headline.split()) < 3:
            fnd.add(idx, "warn", "headline",
                    f"headline looks like a topic label, not a sentence: {headline!r}")

    body_words = 0
    has_source = False
    has_figure = False

    for sh, top in texts:
        tf = sh.text_frame
        is_bottom = top >= BOTTOM_ZONE_IN
        txt = tf.text.strip()
        if txt.lower().startswith("source:"):
            has_source = True
        # per-paragraph and per-run checks
        for para in tf.paragraphs:
            ptext = "".join(r.text for r in para.runs) or para.text
            if BULLET_RE.match(ptext) or (DASH_BULLET_RE.match(ptext) and sh is not headline_shape):
                fnd.add(idx, "error", "bullet",
                        f"bulleted line under a headline: {ptext.strip()[:60]!r}")
            # bullet formatting carried in XML (buChar / buAutoNum)
            pPr = para._pPr
            if pPr is not None and (pPr.find(_qn("a:buChar")) is not None or
                                    pPr.find(_qn("a:buAutoNum")) is not None):
                fnd.add(idx, "warn", "bullet",
                        "paragraph carries a bullet/auto-number format; use spatial layout instead")
            for run in para.runs:
                # The headline is checked like any other text; only the bottom zone
                # (source tags, slide numbers) is exempt from the formatting checks.
                if is_bottom:
                    continue
                if run.font.italic:
                    fnd.add(idx, "error", "format", f"italic text: {run.text.strip()[:40]!r}")
                if run.font.underline:
                    fnd.add(idx, "error", "format", f"underlined text: {run.text.strip()[:40]!r}")
                name = run.font.name
                if name and name.lower() not in ALLOWED_FONTS:
                    fnd.add(idx, "warn", "font",
                            f"font {name!r} is not in the sanctioned set "
                            f"({', '.join(sorted(ALLOWED_FONTS))})")
            # Caps is a formatting check like italics and underline, so it takes the
            # same bottom-zone exemption: source tags carry system and file names
            # ("Source: EPICOR export") that are legitimately upper-case.
            if not is_bottom:
                for m in CAPS_RE.findall(ptext):
                    if m not in CAPS_ALLOW:
                        fnd.add(idx, "warn", "caps", f"all-caps word: {m!r} (avoid all capitals)")
        if sh is not headline_shape and not is_bottom:
            body_words += len(txt.split())
        if DIGIT_RE.search(txt) and sh is not headline_shape and not is_bottom:
            has_figure = True

    # tables and charts carry figures too
    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    if DIGIT_RE.search(cell.text):
                        has_figure = True
        if getattr(sh, "has_chart", False) and sh.has_chart:
            has_figure = True

    # Font-size floors. Nothing checked these, so decks whose prose sat at source-tag size
    # passed stage 6 clean. The floors follow design-tokens.md: 28 pt headline, 18 pt prose
    # body, and a documented exception for dense layouts — 16 pt for two-column/three-row
    # items and 14 pt for table cells and flow-step labels, which are scanned as tokens rather
    # than read as sentences. Runs with no explicit size inherit from the layout and cannot be
    # judged here, so they are skipped rather than guessed at.
    for sh, _top in texts:
        in_table_or_step = getattr(sh, "has_table", False)
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                pt = run.font.size.pt if run.font.size is not None else None
                if pt is None or not run.text.strip():
                    continue
                if sh is headline_shape:
                    if pt < 24:
                        fnd.add(idx, "warn", "fontsize",
                                f"headline set at {pt:g} pt (28 pt is the standard; below 24 pt "
                                f"it stops reading as the assertion)")
                elif pt < 14:
                    fnd.add(idx, "error", "fontsize",
                            f"body text at {pt:g} pt — below the 14 pt floor even for dense "
                            f"layouts; that is source-tag size: {run.text.strip()[:40]!r}")
                elif pt < 18 and not in_table_or_step and len(run.text.split()) > 12:
                    fnd.add(idx, "warn", "fontsize",
                            f"prose run at {pt:g} pt (prose body is 18-24 pt; 14-16 pt is the "
                            f"dense-layout exception for table cells, step labels and column "
                            f"items). Shrinking prose to fit means the slide has too much on "
                            f"it: {run.text.strip()[:40]!r}")

    if body_words > max_words:
        fnd.add(idx, "error", "wordcount",
                f"{body_words} body words (max {max_words}); the claim is too big — split it")
    if has_figure and not has_source:
        fnd.add(idx, "warn", "source",
                "slide carries a figure but has no 'Source:' tag")


def _qn(tag):
    from pptx.oxml.ns import qn
    return qn(tag)


def lint(path, max_words=MAX_BODY_WORDS):
    prs = Presentation(path)
    fnd = Findings()
    for i, slide in enumerate(prs.slides, start=1):
        lint_slide(i, slide, fnd, max_words)
    return fnd


def self_test():
    """Lint a deck built here, asserting the checks reach where the docstring says.

    Regression cases: the headline is subject to the italic/underline/font checks
    (they were once unreachable there, so a shouting, italic, Comic Sans headline
    linted clean), and the bottom zone is exempt from all of them including caps
    (an upper-case system name in a source tag is not a violation).
    """
    import tempfile, os
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def box(x, y, w, h, text, **font):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = text
        for k, v in font.items():
            setattr(run.font, k, v)

    box(0.92, 0.62, 11.52, 0.94, "The backlog doubled in Q3 while inflow stayed flat",
        italic=True, underline=True, name="Comic Sans MS")
    box(0.92, 2.00, 11.52, 2.00, "- one bulleted body line with 42 units", italic=True)
    box(0.92, 6.95, 8.00, 0.40, "Source: EPICOR GENERAL ledger export, 2026-08-19")

    # Slide 2 exists so that every check the docstring claims is exercised actually
    # fires. Before this slide existed the fixture triggered only format/font/bullet,
    # and mutation testing showed the headline, wordcount, source and caps checks
    # could each be deleted outright with `--self-test` still exiting 0 — the gate
    # passed a linter that had stopped linting.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box(0.92, 0.62, 11.52, 1.20,
        "This headline is deliberately far longer than one hundred and ten characters "
        "so that the headline length check has something real to fire on here")
    box(0.92, 2.20, 11.52, 2.60,
        "This body deliberately carries well over the default word budget so the "
        "wordcount check fires: " + " ".join(f"word{i}" for i in range(45))
        + " and the figure 4200 makes this slide carry a number, which with no Source "
          "tag anywhere on the slide is what the source check exists to catch, while "
          "URGENT sits above the bottom zone so the caps check fires too")

    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    # Slide 3 exercises the font-size floors, which nothing checked before: prose shrunk to
    # source-tag size passed stage 6 clean. One run below the hard floor (error), one prose
    # run in the dense-layout band (warn), and one legitimate 14 pt short label that must NOT
    # warn — so the exception is tested as well as the rule.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box(0.92, 0.62, 11.52, 1.20, "Body type below the floor must be an error", size=Pt(28))
    box(0.92, 2.00, 11.52, 0.60, "prose squeezed down to source-tag size to make it fit",
        size=Pt(12))
    box(0.92, 3.00, 11.52, 0.60,
        "a full prose sentence set at sixteen point which is more than twelve words long",
        size=Pt(16))
    box(0.92, 4.00, 3.00, 0.40, "Intake", size=Pt(14))          # legitimate step label

    try:
        prs.save(path)
        found = lint(path).items
    finally:
        os.unlink(path)

    checks = {f"{f['check']}:{f['level']}" for f in found}
    failures = []
    # Assert per CHECK, not per category. Membership on a shared category let one
    # check be deleted invisibly whenever a sibling in the same category still fired.
    for want, why in (("format:error", "italic/underline on the headline must be an error"),
                      ("font:warn", "an unsanctioned headline font must warn"),
                      ("bullet:error", "a dashed body line must be an error"),
                      ("headline:error", "an over-long headline must be an error"),
                      ("wordcount:error", "a body over the word budget must be an error"),
                      ("source:warn", "a figure with no Source: tag must warn"),
                      ("caps:warn", "an all-caps word above the bottom zone must warn"),
                      ("fontsize:error", "body text below the 14 pt floor must be an error"),
                      ("fontsize:warn", "long prose in the dense-layout band must warn")):
        if want not in checks:
            failures.append(f"missing {want} — {why}")
    # The bottom zone stays exempt: slide 1's source tag shouts EPICOR GENERAL and
    # must not be flagged, so caps may fire on slide 2 and must not fire on slide 1.
    caps_s1 = [f for f in found if f["check"] == "caps" and f["slide"] == 1]
    if caps_s1:
        failures.append(
            f"bottom-zone caps must be exempt, got {[c['message'] for c in caps_s1]}")
    # italic and underline both emit check "format", so category membership cannot
    # tell them apart — deleting either one left the other still firing format:error
    # and the mutation went unnoticed. Assert on the distinguishing message text.
    msgs = " | ".join(f["message"] for f in found)
    for frag, why in (("italic text:", "the italic check must fire on italic body text"),
                      ("underlined text:", "the underline check must fire on underlined text")):
        if frag not in msgs:
            failures.append(f"missing {frag!r} — {why}")
    # The dense-layout exception must hold: a short 14 pt label is legitimate and must not
    # be flagged, or the check just re-imposes the blanket rule the docs retired.
    if any(f["check"] == "fontsize" and "Intake" in f["message"] for f in found):
        failures.append("a short 14 pt step label must NOT be flagged (dense-layout exception)")

    if failures:
        for f in failures:
            print("FAIL " + f, file=sys.stderr)
        return 1
    print(f"self-test passed ({len(found)} findings, all as expected)")
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(description="Lint a .pptx against the assertion-evidence checklist.")
    ap.add_argument("pptx", nargs="?", help="path to the .pptx to audit")
    ap.add_argument("--self-test", action="store_true",
                    help="lint a deck built here that exercises each check, and exit")
    ap.add_argument("--json", action="store_true", help="machine-readable output")
    ap.add_argument("--max-body-words", type=int, default=MAX_BODY_WORDS)
    args = ap.parse_args(argv)

    if args.self_test:
        return self_test()
    if not args.pptx:
        ap.error("a .pptx path is required (or use --self-test)")

    fnd = lint(args.pptx, args.max_body_words)

    if args.json:
        print(json.dumps({"findings": fnd.items,
                          "errors": len(fnd.errors),
                          "total": len(fnd.items)}, indent=2))
    else:
        if not fnd.items:
            print("clean — no assertion-evidence violations found")
        for f in fnd.items:
            print(f"[{f['level'].upper():5}] slide {f['slide']:>2} {f['check']:<9} {f['message']}")
        print(f"\n{len(fnd.errors)} error(s), {len(fnd.items) - len(fnd.errors)} warning(s)")

    return 1 if fnd.errors else 0


if __name__ == "__main__":
    sys.exit(main())
